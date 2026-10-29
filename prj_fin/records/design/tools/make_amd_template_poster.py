# Build an AMD-template-style poster for the Zynq-7020 video node.
# Output is 13.333 x 7.5 in (16:9), matching AMD_poster_template.pptx.
from pathlib import Path
import sys
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(r"D:/Xilinx/Prj/FIN")
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "records/_poster_tmp/poster_amd_template.pptx"
ASSETS = ROOT / "records/design/assets"

INK = RGBColor(32, 32, 32)
GRAY = RGBColor(92, 92, 92)
PURPLE = RGBColor(122, 45, 94)
CYAN = RGBColor(0, 150, 150)
LIGHT = RGBColor(244, 242, 240)
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(28, 28, 28)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])


def text(x, y, w, h, runs, size=10, bold=False, font="Times New Roman",
         color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = spacing
    if isinstance(runs, str):
        runs = [(runs, bold, size, color)]
    for item in runs:
        t, b, s, c = item
        r = p.add_run()
        r.text = t
        r.font.name = font
        r.font.size = Pt(s)
        r.font.bold = b
        r.font.color.rgb = c
    return box


def rich(x, y, w, h, paragraphs, font="Times New Roman", size=9,
         color=INK, spacing=1.0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for para in paragraphs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(2)
        p.line_spacing = spacing
        for t, b, s, c in para:
            r = p.add_run()
            r.text = t
            r.font.name = font
            r.font.size = Pt(s)
            r.font.bold = b
            r.font.color.rgb = c
    return box


def bar(x, y, w, h, label):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = BLACK
    shp.line.fill.background()
    tf = shp.text_frame
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.08)
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = label
    r.font.name = "Arial Black"
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = WHITE
    return shp


def card(x, y, w, h, title, body):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = 0.08
    shp.fill.solid()
    shp.fill.fore_color.rgb = WHITE
    shp.line.color.rgb = RGBColor(170, 170, 170)
    shp.line.width = Pt(0.9)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.16)
    tf.margin_right = Inches(0.14)
    tf.margin_top = Inches(0.08)
    tf.margin_bottom = Inches(0.06)
    p = tf.paragraphs[0]
    p.space_after = Pt(3)
    r = p.add_run()
    r.text = title
    r.font.name = "Arial"
    r.font.size = Pt(10.5)
    r.font.bold = True
    r.font.color.rgb = PURPLE
    p2 = tf.add_paragraph()
    p2.line_spacing = 1.0
    r2 = p2.add_run()
    r2.text = body
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = INK
    return shp


def pic(path, x, y, w, h):
    im = Image.open(path)
    iw, ih = im.size
    box_ratio = w / h
    img_ratio = iw / ih
    if img_ratio > box_ratio:
        nw, nh = w, w / img_ratio
    else:
        nh, nw = h, h * img_ratio
    nx, ny = x + (w - nw) / 2, y + (h - nh) / 2
    return slide.shapes.add_picture(str(path), Inches(nx), Inches(ny), Inches(nw), Inches(nh))


def rule(x, y, w, color=RGBColor(190, 190, 190), width=0.8):
    shp = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
    shp.line.color.rgb = color
    shp.line.width = Pt(width)
    return shp


# ---- top ----
text(0.28, 0.08, 3.2, 0.18, "(c) 2026 Uie_Team  |  AMD FPGA Innovation Design Competition", size=7.5,
     color=GRAY, font="Arial")
text(0.42, 0.33, 8.6, 0.55, "ONE FRAME. TWO TRUTHS.", size=31, bold=True,
     font="Times New Roman", color=INK)
text(0.44, 0.88, 8.5, 0.30, "A full-hardware video node on Zynq-7020", size=13.5, bold=False,
     font="Times New Roman", color=GRAY)
text(0.44, 1.16, 8.3, 0.25, "多源视频同帧对照处理系统", size=11, bold=False,
     font="Microsoft YaHei", color=PURPLE)
text(0.44, 1.43, 8.3, 0.38,
     [("Team: ", True, 10.5, INK), ("Uie_Team  ·  Hunan University of Technology (湖南工业大学), Hunan Province", False, 10.5, INK)],
     font="Times New Roman")

# school emblem + name (top right)
pic(ASSETS / "hut_emblem.png", 8.98, 0.10, 0.70, 0.72)
text(9.70, 0.12, 1.36, 0.26, "湖南工业大学", size=9.5, bold=True, font="Microsoft YaHei", color=PURPLE)
text(9.70, 0.40, 1.36, 0.34, "HUNAN UNIVERSITY\nOF TECHNOLOGY", size=5.6, bold=True,
     font="Arial", color=PURPLE)
# AMD logo
pic(ASSETS / "amd_logo_black.png", 8.98, 0.92, 1.25, 0.30)
text(8.98, 1.25, 1.25, 0.20, "ZYNQ-7020", size=7.5, bold=True, font="Arial", color=GRAY)

# board photo
pic(ASSETS / "poster_board.jpg", 11.15, 0.12, 1.90, 1.43)
text(11.10, 1.62, 1.95, 0.31, "RK-ZYNQ7020-F V1.0\nxc7z020clg484-2 · 1 GbE", size=7.2,
     color=GRAY, font="Arial")

rule(0.30, 1.88, 12.75, color=RGBColor(60, 60, 60), width=1.2)

# ---- left column ----
bar(0.30, 2.02, 3.78, 0.38, "INTRODUCTION")
text(0.34, 2.48, 3.70, 1.55,
     "Remote video nodes usually have one Ethernet cable and no PC. The operator must know whether the frame arrived, whether processing took effect, and how old the picture is. This node answers all three from one frame: 512x300 RGB565 enters RGMII, is reassembled and committed to DDR only after three gates, then read back once for both panes. The left pane is the original; the right pane is the same frame after geometry and a fixed 15-cycle effect chain. The seam is a per-pixel decision, not two screenshots.",
     size=9.2, font="Times New Roman", color=INK, spacing=1.0)

pic(ASSETS / "poster_frame_split.png", 0.45, 4.00, 3.35, 1.96)
text(0.34, 5.99, 3.70, 0.26,
     "Same 512x300 frame  ·  seam 50/50  ·  right pane 0 deg + Sobel",
     size=7.2, color=GRAY, font="Arial")

text(0.34, 6.27, 3.70, 0.20, "Key technical issues addressed:", size=9.0, bold=True,
     font="Times New Roman", color=PURPLE)
text(0.34, 6.48, 3.70, 0.64,
     "1. 1 GbE UDP: offset reassembly; commit = bytes + 300 rows + no bad packet.\n"
     "2. Same-frame A/B: one coordinate stream, one row map, 15-cycle chain.\n"
     "3. Telemetry: aborts, rows_max, gaps, pane_bad; drop 1/100 reddens it.",
     size=6.7, font="Times New Roman", color=INK, spacing=0.90)

# ---- middle column ----
bar(4.30, 2.02, 4.15, 0.38, "SYSTEM ARCHITECTURE")
pic(ASSETS / "poster_arch.png", 4.30, 2.48, 4.15, 1.87)
text(4.30, 4.38, 4.15, 0.24,
     "RGMII -> UDP -> DDR ping-pong -> geometry -> effects -> A/B seam -> HDMI",
     size=7.0, color=GRAY, font="Arial")

card(4.30, 4.80, 4.15, 1.10, "THREE-GATE COMMIT",
     "A frame is committed only when 307200 bytes arrived, all 300 source rows were written, and no packet failed. Otherwise the last good frame stays on screen and aborts + rows_max expose the reason.")
card(4.30, 6.02, 4.15, 1.10, "SAME-FRAME A/B",
     "Original and processed panes share one source-coordinate stream, one row map and one 15-cycle effect chain. The seam is a per-pixel select; rotation or stacked effects cannot make the panes disagree.")

# ---- right column ----
bar(8.65, 2.02, 4.40, 0.38, "RESULT")
text(8.70, 2.48, 4.30, 0.55,
     "RGMII -> HDMI entirely in PL; PS only control, SD playback and telemetry. 1024x600 at 59.5 Hz, 1 GbE UDP ingest.",
     size=9.2, font="Times New Roman", color=INK)

rows = [
    ("Setup WNS", "1.341 ns  ·  0 failing endpoints"),
    ("Hold WHS", "0.025 ns  ·  positive margin"),
    ("BRAM", "70.0%  ·  98 / 140 tiles"),
    ("8 s stream", "241 frames  ·  29.95 fps"),
    ("UART battery", "105 commands  ·  PASS"),
    ("QSPI cold boot", "DONE / EOS  ·  no CRC error"),
    ("Fault injection", "22.0 / 0 / 18.9 fps  ·  PASS"),
]
y0 = 3.15
for i, (k, v) in enumerate(rows):
    yy = y0 + i * 0.30
    fill = RGBColor(248, 248, 248) if i % 2 == 0 else WHITE
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.70), Inches(yy), Inches(4.30), Inches(0.28))
    shp.fill.solid(); shp.fill.fore_color.rgb = fill; shp.line.fill.background()
    text(8.76, yy + 0.04, 1.35, 0.20, k, size=8.2, bold=True, font="Arial", color=GRAY)
    text(10.10, yy + 0.04, 2.82, 0.20, v, size=8.6, bold=False, font="Arial", color=INK)

text(8.70, 5.48, 4.30, 0.22, "OBSERVABILITY", size=10.5, bold=True, font="Arial Black", color=PURPLE)
text(8.70, 5.74, 4.30, 0.68,
     "link / 1 Hz UDP telemetry: aborts, rows_max, stall, gap_min/max and pane_bad. The same counters are readable over UART, JTAG and the wire.",
     size=8.6, font="Times New Roman", color=INK)
# fault badge
shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.70), Inches(6.55), Inches(4.30), Inches(0.42))
shp.adjustments[0] = 0.15
shp.fill.solid(); shp.fill.fore_color.rgb = RGBColor(245, 236, 241)
shp.line.color.rgb = PURPLE; shp.line.width = Pt(0.9)
tf = shp.text_frame; tf.margin_left = Inches(0.12); tf.margin_top = 0; tf.margin_bottom = 0
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "DROP 1 IN 100  ->  ABORTS RISE  ->  RECOVERY"
r.font.name = "Arial"; r.font.size = Pt(9); r.font.bold = True; r.font.color.rgb = PURPLE

# footer
rule(0.30, 7.18, 12.75, color=RGBColor(190, 190, 190), width=0.7)
text(0.30, 7.23, 9.8, 0.18,
     "AMD, the AMD Arrow logo and combinations thereof are trademarks of Advanced Micro Devices, Inc. Zynq is a trademark of AMD. All other trademarks are the property of their respective owners.",
     size=6.2, color=GRAY, font="Arial")
pic(ASSETS / "amd_logo_black.png", 12.35, 7.10, 0.68, 0.18)

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(OUT)
