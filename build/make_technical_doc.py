# -*- coding: utf-8 -*-
"""Render report/design-report.md into technical_document.docx.

No local template, no absolute paths. Images are resolved relative to the
repository root; the markdown stays the single source of truth.
"""
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "report", "design-report.md")
OUT = os.path.join(ROOT, "technical_document.docx")

BODY_CN = "宋体"
HEAD_CN = "黑体"
BODY_EN = "Times New Roman"
MONO = "Consolas"
DARK = RGBColor(0x1F, 0x3A, 0x5F)
GRAY = RGBColor(0x59, 0x59, 0x59)


def set_run_font(run, size=10.5, bold=False, color=None, cn=BODY_CN, en=BODY_EN,
                 italic=False, mono=False):
    run.font.name = MONO if mono else en
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), cn)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_para(p, indent=True, before=0, after=4, line=1.35, align=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if indent:
        pf.first_line_indent = Pt(21)
    if align is not None:
        p.alignment = align


def add_runs(p, text, size=10.5, color=None, cn=BODY_CN, en=BODY_EN):
    """Handle **bold** and `code` in one line."""
    pos = 0
    for m in re.finditer(r"(\*\*.*?\*\*|`[^`]*`)", text):
        if m.start() > pos:
            r = p.add_run(text[pos:m.start()])
            set_run_font(r, size=size, color=color, cn=cn, en=en)
        tok = m.group(0)
        if tok.startswith("**"):
            r = p.add_run(tok[2:-2])
            set_run_font(r, size=size, bold=True, color=color, cn=cn, en=en)
        else:
            r = p.add_run(tok[1:-1])
            set_run_font(r, size=size - 0.5, color=color, cn=cn, en=en, mono=True)
        pos = m.end()
    if pos < len(text):
        r = p.add_run(text[pos:])
        set_run_font(r, size=size, color=color, cn=cn, en=en)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_cell_margins(table, top=40, bottom=40, left=70, right=70):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement("w:" + side)
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "BFBFBF")
        borders.append(el)
    tblPr.append(borders)


def add_table(doc, rows):
    ncol = max(len(r) for r in rows)
    rows = [r + [""] * (ncol - len(r)) for r in rows]
    table = doc.add_table(rows=0, cols=ncol)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    set_cell_margins(table)
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j, val in enumerate(row):
            cells[j].text = ""
            p = cells[j].paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (i == 0 or j > 0) else WD_ALIGN_PARAGRAPH.LEFT
            add_runs(p, val, size=9, color=None)
            if i == 0:
                shade(cells[j], "E8EEF7")
                for r in p.runs:
                    r.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_code(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    set_cell_margins(table, top=60, bottom=60, left=90, right=90)
    cell = table.rows[0].cells[0]
    shade(cell, "F4F5F7")
    cell.text = ""
    first = True
    for line in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(line if line else " ")
        set_run_font(r, size=8.5, color=RGBColor(0x22, 0x22, 0x22), mono=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_figure(doc, path, caption, width_cm=14.4):
    full = os.path.join(ROOT, path)
    if not os.path.exists(full):
        full = os.path.join(ROOT, "report", path)
    if not os.path.exists(full):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("[missing figure: %s]" % path)
        set_run_font(r, size=9, color=GRAY)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(full, width=Cm(width_cm))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(6)
    add_runs(c, caption, size=9, color=GRAY)


def add_heading(doc, text, level):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    if level == 1:
        set_para(p, indent=False, before=12, after=5, line=1.15)
        r = p.add_run(text)
        set_run_font(r, size=15, bold=True, color=DARK, cn=HEAD_CN, en="Arial")
    elif level == 2:
        set_para(p, indent=False, before=9, after=3, line=1.15)
        r = p.add_run(text)
        set_run_font(r, size=12.5, bold=True, color=DARK, cn=HEAD_CN, en="Arial")
    else:
        set_para(p, indent=False, before=6, after=2, line=1.15)
        r = p.add_run(text)
        set_run_font(r, size=11, bold=True, color=RGBColor(0x33, 0x33, 0x33), cn=HEAD_CN, en="Arial")
    return p


def split_table_row(line):
    # Protect escaped pipes inside cells.
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    parts = re.split(r"(?<!\\)\|", line)
    return [p.strip().replace("\\|", "|") for p in parts]


def is_sep_row(line):
    cells = split_table_row(line)
    return all(re.fullmatch(r":?-{2,}:?", c.replace(" ", "")) for c in cells if c != "")


def parse_and_render(doc, lines):
    i = 0
    n = len(lines)
    title_seen = False
    subtitle_seen = False
    while i < n:
        raw = lines[i]
        line = raw.strip()
        if not line:
            i += 1
            continue
        if line.startswith("# "):
            if not title_seen:
                title_seen = True
                title = line[2:].strip()
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(60)
                p.paragraph_format.space_after = Pt(6)
                r = p.add_run(title)
                set_run_font(r, size=22, bold=True, color=DARK, cn=HEAD_CN, en="Arial")
            i += 1
            continue
        if line.startswith("## "):
            text = line[3:].strip()
            if title_seen and not subtitle_seen:
                subtitle_seen = True
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(4)
                r = p.add_run(text)
                set_run_font(r, size=13, color=RGBColor(0x33, 0x33, 0x33), cn=HEAD_CN, en="Arial")
                p2 = doc.add_paragraph()
                p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p2.paragraph_format.space_after = Pt(18)
                r2 = p2.add_run("AMD FPGA 创新设计大赛 · 初级组 · 自命题赛道")
                set_run_font(r2, size=9.5, color=GRAY)
                i += 1
                continue
            add_heading(doc, text, 1)
            i += 1
            continue
        if line.startswith("### "):
            add_heading(doc, line[4:].strip(), 2)
            i += 1
            continue
        if line.startswith("#### "):
            add_heading(doc, line[5:].strip(), 3)
            i += 1
            continue
        if line == "---":
            doc.add_page_break()
            i += 1
            continue
        if line.startswith("```"):
            i += 1
            code = []
            while i < n and not lines[i].strip().startswith("```"):
                code.append(lines[i].rstrip("\n"))
                i += 1
            if i < n:
                i += 1
            add_code(doc, code)
            continue
        m = re.match(r"!\[(.*?)\]\((.*?)\)", line)
        if m:
            add_figure(doc, m.group(2), m.group(1))
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rowline = lines[i].strip()
                if not is_sep_row(rowline):
                    rows.append(split_table_row(rowline))
                i += 1
            if rows:
                add_table(doc, rows)
            continue
        if line.startswith("- "):
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Cm(0.74)
            pf.first_line_indent = Cm(-0.37)
            pf.space_before = Pt(0)
            pf.space_after = Pt(2)
            pf.line_spacing = 1.3
            add_runs(p, "• " + line[2:].strip())
            i += 1
            continue
        if re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.left_indent = Cm(0.74)
            pf.first_line_indent = Cm(-0.37)
            pf.space_before = Pt(0)
            pf.space_after = Pt(2)
            pf.line_spacing = 1.3
            add_runs(p, line)
            i += 1
            continue
        if line.startswith("> "):
            p = doc.add_paragraph()
            set_para(p, indent=False, before=2, after=2, line=1.25)
            p.paragraph_format.left_indent = Cm(0.6)
            add_runs(p, line[2:].strip(), color=GRAY)
            i += 1
            continue
        p = doc.add_paragraph()
        set_para(p)
        add_runs(p, line)
        i += 1


def add_field(paragraph, instr):
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instr)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)


def main():
    if not os.path.exists(SRC):
        print("missing %s" % SRC)
        return 1
    with open(SRC, encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.3)
    sec.right_margin = Cm(2.3)

    style = doc.styles["Normal"]
    style.font.name = BODY_EN
    style.font.size = Pt(10.5)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), BODY_CN)

    # Header
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("多源视频流多效果变换处理系统")
    set_run_font(hr, size=8.5, color=GRAY, cn=BODY_CN, en=BODY_EN)
    pPr = hp._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "BFBFBF")
    pBdr.append(bottom)
    pPr.append(pBdr)

    # Footer with page number
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("第 ")
    set_run_font(r, size=9, color=GRAY)
    add_field(fp, "PAGE")
    r = fp.add_run(" 页")
    set_run_font(r, size=9, color=GRAY)

    parse_and_render(doc, lines)

    # Document properties
    doc.core_properties.title = "多源视频流多效果变换处理系统"
    doc.core_properties.subject = "基于 Zynq-7020 的全硬件视频通路与可观测性设计"
    doc.core_properties.author = ""

    doc.save(OUT)
    print("WROTE %s" % OUT)
    return 0


if __name__ == "__main__":
    sys.exit(main())
